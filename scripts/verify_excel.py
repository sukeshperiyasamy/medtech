import openpyxl

wb = openpyxl.load_workbook("IIT_Jodhpur_MedTech_Students_Directory_Master.xlsx", data_only=False)
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    formulas = []
    values = []
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if cell.value and str(cell.value).startswith("="):
                formulas.append((cell.coordinate, str(cell.value)))
    print(f"Sheet '{sheetname}': {ws.max_row} rows, {ws.max_column} columns, {len(formulas)} formulas")
    if formulas:
        print("  Sample formulas:", formulas[:5])

# Also check data_only=True to see if anything precomputed or any string errors
wb_data = openpyxl.load_workbook("IIT_Jodhpur_MedTech_Students_Directory_Master.xlsx", data_only=True)
err_list = ["#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"]
found_errs = []
for sheetname in wb_data.sheetnames:
    ws = wb_data[sheetname]
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if str(cell.value) in err_list:
                found_errs.append((sheetname, cell.coordinate, cell.value))

print(f"Direct formula string errors found: {len(found_errs)} -> {found_errs}")
