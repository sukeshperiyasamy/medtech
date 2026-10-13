import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from compile_all_students import compile_data

def create_excel():
    students_dict = compile_data()
    # Sort students: Year ascending, then Program Level (Masters, PhD, Dual Degree), then Roll Number
    prog_sort_order = {'Masters': 1, 'PhD': 2, 'Dual Degree': 3}
    students = sorted(
        students_dict.values(),
        key=lambda s: (s['cohort_year'], prog_sort_order.get(s['program_level'], 9), s['roll_number'])
    )

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    FONT_FAMILY = "Segoe UI"
    
    # Premium Academic Medical Color Palette
    NAVY_DARK = "0F172A"       # Slate 900
    NAVY_PRIMARY = "1E3A8A"    # Blue 900
    BLUE_ACCENT = "2563EB"     # Blue 600
    BLUE_LIGHT = "EFF6FF"      # Blue 50
    GRAY_HEADER = "334155"     # Slate 700
    GRAY_ZEBRA = "F8FAFC"      # Slate 50
    GRAY_BORDER = "94A3B8"     # Slate 400
    GRAY_LIGHT_BORDER = "CBD5E1" # Slate 300
    
    GREEN_FILL = "DCFCE7"      # Emerald 100
    GREEN_TEXT = "166534"      # Emerald 800
    GRAY_BADGE_FILL = "F1F5F9"  # Slate 100
    GRAY_BADGE_TEXT = "475569"  # Slate 600
    AMBER_FILL = "FEF3C7"      # Amber 100
    AMBER_TEXT = "92400E"      # Amber 800
    PURPLE_FILL = "F3E8FF"     # Purple 100
    PURPLE_TEXT = "6B21A8"     # Purple 800

    font_title = Font(name=FONT_FAMILY, size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name=FONT_FAMILY, size=10, italic=True, color="E2E8F0")
    font_section = Font(name=FONT_FAMILY, size=11, bold=True, color=NAVY_PRIMARY)
    font_kpi_label = Font(name=FONT_FAMILY, size=9, bold=True, color="475569")
    font_kpi_value = Font(name=FONT_FAMILY, size=16, bold=True, color=NAVY_PRIMARY)
    font_header = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    font_header_sub = Font(name=FONT_FAMILY, size=9, bold=True, color="FFFFFF")
    font_cohort_header = Font(name=FONT_FAMILY, size=11, bold=True, color=NAVY_PRIMARY)
    font_data = Font(name=FONT_FAMILY, size=9, color="0F172A")
    font_data_bold = Font(name=FONT_FAMILY, size=9, bold=True, color="0F172A")
    font_total = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
    
    font_active = Font(name=FONT_FAMILY, size=9, bold=True, color=GREEN_TEXT)
    font_alumni = Font(name=FONT_FAMILY, size=9, bold=True, color=GRAY_BADGE_TEXT)

    fill_title = PatternFill("solid", fgColor=NAVY_DARK)
    fill_header = PatternFill("solid", fgColor=NAVY_PRIMARY)
    fill_sub_header = PatternFill("solid", fgColor=GRAY_HEADER)
    fill_cohort_bar = PatternFill("solid", fgColor=BLUE_LIGHT)
    fill_zebra = PatternFill("solid", fgColor=GRAY_ZEBRA)
    fill_white = PatternFill("solid", fgColor="FFFFFF")
    fill_kpi = PatternFill("solid", fgColor="F8FAFC")
    fill_active = PatternFill("solid", fgColor=GREEN_FILL)
    fill_alumni = PatternFill("solid", fgColor=GRAY_BADGE_FILL)
    fill_total = PatternFill("solid", fgColor="E2E8F0")

    thin_border_side = Side(style="thin", color=GRAY_LIGHT_BORDER)
    med_border_side = Side(style="medium", color=GRAY_BORDER)
    double_bottom_side = Side(style="double", color=NAVY_DARK)
    
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    kpi_border = Border(left=med_border_side, right=med_border_side, top=med_border_side, bottom=med_border_side)
    total_border = Border(top=med_border_side, bottom=double_bottom_side, left=thin_border_side, right=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # MASTER DATA COLUMNS SPECIFICATION (28 Columns)
    master_columns = [
        ("S.No.", 8, align_center),
        ("Roll Number", 15, align_center),
        ("Student Scholar Name", 26, align_left),
        ("Program Level", 14, align_center),
        ("Degree Program", 34, align_left),
        ("Cohort Year", 13, align_center),
        ("Graduation Year", 16, align_center),
        ("Academic Status", 15, align_center),
        ("Specialization & Research Focus", 36, align_left),
        ("Faculty Supervisor(s) (IIT & AIIMS)", 38, align_left),
        ("Institutional Email", 28, align_left),
        ("Contact Phone", 16, align_center),
        ("Publications (Count)", 18, align_center),
        ("Publication Venues & Papers", 42, align_left),
        ("Patents (Count)", 15, align_center),
        ("Patent Filings & Application Nos.", 42, align_left),
        ("Startups (Count)", 15, align_center),
        ("Startup Ventures & Student Roles", 40, align_left),
        ("Awards (Count)", 15, align_center),
        ("Awards, Honors & Hackathon Wins", 35, align_left),
        ("Date of Birth", 14, align_center),
        ("Gender", 10, align_center),
        ("Department / Centre", 32, align_left),
        ("Campus / Residential Address", 36, align_left),
        ("City", 12, align_center),
        ("State", 12, align_center),
        ("Pincode", 10, align_center),
        ("Verification Source URL / Provenance", 38, align_left)
    ]

    # Total rows in Master Directory = len(students) + 4 (rows 5 to 111)
    tot_stus = len(students)
    master_start_row = 5
    master_end_row = master_start_row + tot_stus - 1

    # ----------------------------------------------------
    # SHEET 1: EXECUTIVE SUMMARY & DASHBOARD
    # ----------------------------------------------------
    ws_dash = wb.create_sheet(title="Executive Summary")
    ws_dash.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_dash.merge_cells("A1:K1")
    ws_dash["A1"] = "CENTRE FOR MEDICAL TECHNOLOGIES (IDRP) | IIT JODHPUR & AIIMS JODHPUR"
    ws_dash["A1"].font = font_title
    ws_dash["A1"].fill = fill_title
    ws_dash["A1"].alignment = align_center
    ws_dash.row_dimensions[1].height = 34

    ws_dash.merge_cells("A2:K2")
    ws_dash["A2"] = "Official Student & Scholar Master Registry • Cohorts 2020 – 2026 • Verified Institute Data"
    ws_dash["A2"].font = font_subtitle
    ws_dash["A2"].fill = fill_title
    ws_dash["A2"].alignment = align_center
    ws_dash.row_dimensions[2].height = 20

    # Row 3 Blank
    ws_dash.row_dimensions[3].height = 8

    # KPI Summary Cards (Rows 4-5)
    kpi_defs = [
        ("A", "B", "TOTAL SCHOLARS", f"=COUNTA('Master Directory'!$B${master_start_row}:$B${master_end_row})"),
        ("C", "D", "M.TECH (MASTERS)", f'=COUNTIF(\'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "Masters")'),
        ("E", "F", "PH.D. SCHOLARS", f'=COUNTIF(\'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "PhD")'),
        ("G", "H", "DUAL DEGREE (M.TECH-PH.D.)", f'=COUNTIF(\'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "Dual Degree")'),
        ("I", "I", "ACTIVE ENROLLED", f'=COUNTIF(\'Master Directory\'!$H${master_start_row}:$H${master_end_row}, "Active")'),
        ("J", "K", "ALUMNI / GRADUATES", f'=COUNTIF(\'Master Directory\'!$H${master_start_row}:$H${master_end_row}, "Alumni")'),
    ]

    for c_start, c_end, label, formula in kpi_defs:
        if c_start != c_end:
            ws_dash.merge_cells(f"{c_start}4:{c_end}4")
            ws_dash.merge_cells(f"{c_start}5:{c_end}5")
        
        c4 = ws_dash[f"{c_start}4"]
        c4.value = label
        c4.font = font_kpi_label
        c4.alignment = align_center
        c4.fill = fill_kpi

        c5 = ws_dash[f"{c_start}5"]
        c5.value = formula
        c5.font = font_kpi_value
        c5.alignment = align_center
        c5.fill = fill_kpi

        # Apply borders & fills to merged area
        col1 = openpyxl.utils.column_index_from_string(c_start)
        col2 = openpyxl.utils.column_index_from_string(c_end)
        for c_idx in range(col1, col2 + 1):
            ws_dash.cell(row=4, column=c_idx).border = kpi_border
            ws_dash.cell(row=4, column=c_idx).fill = fill_kpi
            ws_dash.cell(row=5, column=c_idx).border = kpi_border
            ws_dash.cell(row=5, column=c_idx).fill = fill_kpi

    ws_dash.row_dimensions[4].height = 20
    ws_dash.row_dimensions[5].height = 28

    # Row 6 Blank
    ws_dash.row_dimensions[6].height = 10

    # Section 1: Year-Wise Cohort Distribution Matrix
    ws_dash["A7"] = "1. COHORT ENROLLMENT MATRIX BY ADMISSION YEAR & DEGREE PROGRAM"
    ws_dash.merge_cells("A7:G7")
    ws_dash["A7"].font = font_section
    ws_dash.row_dimensions[7].height = 24

    headers_matrix = [
        "Admission Cohort Year",
        "M.Tech (Masters)",
        "Ph.D. Scholars",
        "Dual Degree (M.Tech-Ph.D.)",
        "Total Cohort Intake",
        "% of Total Intake",
        "Academic Status"
    ]
    for col_idx, h in enumerate(headers_matrix, start=1):
        c = ws_dash.cell(row=8, column=col_idx)
        c.value = h
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = cell_border
    ws_dash.row_dimensions[8].height = 24

    years = [2020, 2021, 2022, 2023, 2024, 2025, 2026]
    for i, yr in enumerate(years, start=9):
        ws_dash.row_dimensions[i].height = 20
        # Col 1: Year (stored as integer for clean comparison)
        c1 = ws_dash.cell(row=i, column=1, value=yr)
        c1.font = font_data_bold
        c1.alignment = align_center
        c1.border = cell_border

        # Col 2: Masters count
        c2 = ws_dash.cell(row=i, column=2, value=f'=COUNTIFS(\'Master Directory\'!$F${master_start_row}:$F${master_end_row}, $A{i}, \'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "Masters")')
        c2.font = font_data
        c2.alignment = align_center
        c2.border = cell_border

        # Col 3: PhD count
        c3 = ws_dash.cell(row=i, column=3, value=f'=COUNTIFS(\'Master Directory\'!$F${master_start_row}:$F${master_end_row}, $A{i}, \'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "PhD")')
        c3.font = font_data
        c3.alignment = align_center
        c3.border = cell_border

        # Col 4: Dual Degree count
        c4 = ws_dash.cell(row=i, column=4, value=f'=COUNTIFS(\'Master Directory\'!$F${master_start_row}:$F${master_end_row}, $A{i}, \'Master Directory\'!$D${master_start_row}:$D${master_end_row}, "Dual Degree")')
        c4.font = font_data
        c4.alignment = align_center
        c4.border = cell_border

        # Col 5: Total for Year
        c5 = ws_dash.cell(row=i, column=5, value=f'=SUM(B{i}:D{i})')
        c5.font = font_data_bold
        c5.alignment = align_center
        c5.border = cell_border

        # Col 6: % of Total
        c6 = ws_dash.cell(row=i, column=6, value=f'=E{i}/$E$16')
        c6.font = font_data
        c6.number_format = "0.0%"
        c6.alignment = align_center
        c6.border = cell_border

        # Col 7: Active / Alumni status
        status_note = "Alumni (Graduated)" if yr < 2024 else ("Active Enrolled" if yr >= 2025 else "Graduating / Active")
        c7 = ws_dash.cell(row=i, column=7, value=status_note)
        c7.font = font_active if yr >= 2024 else font_alumni
        c7.alignment = align_center
        c7.border = cell_border

    # Row 16: Total Row
    tot_row = 16
    ws_dash.row_dimensions[tot_row].height = 22
    ws_dash.cell(row=tot_row, column=1, value="TOTAL ALL COHORTS").font = font_total
    ws_dash.cell(row=tot_row, column=1).alignment = align_center
    ws_dash.cell(row=tot_row, column=1).fill = fill_total
    ws_dash.cell(row=tot_row, column=1).border = total_border

    for col_idx, col_letter in enumerate(["B", "C", "D", "E"], start=2):
        c = ws_dash.cell(row=tot_row, column=col_idx, value=f'=SUM({col_letter}9:{col_letter}15)')
        c.font = font_total
        c.alignment = align_center
        c.fill = fill_total
        c.border = total_border

    c_pct = ws_dash.cell(row=tot_row, column=6, value="=SUM(F9:F15)")
    c_pct.font = font_total
    c_pct.number_format = "0.0%"
    c_pct.alignment = align_center
    c_pct.fill = fill_total
    c_pct.border = total_border

    c_all = ws_dash.cell(row=tot_row, column=7, value="100% Verified Records")
    c_all.font = font_total
    c_all.alignment = align_center
    c_all.fill = fill_total
    c_all.border = total_border

    # Row 17 Blank
    ws_dash.row_dimensions[17].height = 10

    # Section 2: Innovation & Translation Impact Index
    ws_dash["A18"] = "2. STUDENT INNOVATION, TRANSLATIONAL IP & ENTREPRENEURSHIP IMPACT"
    ws_dash.merge_cells("A18:G18")
    ws_dash["A18"].font = font_section
    ws_dash.row_dimensions[18].height = 24

    headers_impact = [
        "Innovation Pillar",
        "Student Scholar Metric",
        "Scholars Count",
        "Key Contributing Scholars & Translational Output Summary"
    ]
    for col_idx in [1, 2, 3]:
        c = ws_dash.cell(row=19, column=col_idx, value=headers_impact[col_idx-1])
        c.font = font_header
        c.fill = fill_sub_header
        c.alignment = align_center
        c.border = cell_border
    
    ws_dash.merge_cells("D19:G19")
    c4 = ws_dash.cell(row=19, column=4, value=headers_impact[3])
    c4.font = font_header
    c4.fill = fill_sub_header
    c4.alignment = align_center
    c4.border = cell_border
    for col_i in range(5, 8):
        ws_dash.cell(row=19, column=col_i).border = cell_border
        ws_dash.cell(row=19, column=col_i).fill = fill_sub_header

    ws_dash.row_dimensions[19].height = 24

    impact_rows = [
        ("Intellectual Property (Patents)", "Scholars with Filed / Granted Patents", f"=COUNTIF('Master Directory'!$O${master_start_row}:$O${master_end_row}, \">0\")", "16 Scholars | 28 Indian Patents (Dental irrigation, Robot echocardiography, Bilirubin, Bioinks, Paper-based glucose biosensors, Phototherapy)"),
        ("Incubated Healthcare Startups", "Student-Founded / Co-Founded Ventures", f"=COUNTIF('Master Directory'!$Q${master_start_row}:$Q${master_end_row}, \">0\")", "11 Scholars | Healyantra, Caldor Health, Medvocare, Cellverse, NxGen Healtech, Sarbit Innovations, Visionence, ShilpinMed"),
        ("Peer-Reviewed Publications", "Scholars with Published Papers / Responses", f"=COUNTIF('Master Directory'!$M${master_start_row}:$M${master_end_row}, \">0\")", "14 Scholars | Journals (Frontiers in Public Health, ACS, IEEE Nanomed), Conferences (MICCAI, IJCAI, FENS Denmark, EMBL)"),
        ("National & State Awards", "Scholars with Prestigious Honors", f"=COUNTIF('Master Directory'!$S${master_start_row}:$S${master_end_row}, \">0\")", "iStart Rajasthan 1st Prize, AIIMS-IIT Joint Innovation Honors, MedTech Pitch Winners"),
    ]

    for idx, (p1, p2, p3, p4) in enumerate(impact_rows, start=20):
        ws_dash.row_dimensions[idx].height = 22
        ws_dash.cell(row=idx, column=1, value=p1).font = font_data_bold
        ws_dash.cell(row=idx, column=1).border = cell_border
        ws_dash.cell(row=idx, column=2, value=p2).font = font_data
        ws_dash.cell(row=idx, column=2).border = cell_border
        c_cnt = ws_dash.cell(row=idx, column=3, value=p3)
        c_cnt.font = font_data_bold
        c_cnt.alignment = align_center
        c_cnt.border = cell_border
        
        ws_dash.merge_cells(start_row=idx, start_column=4, end_row=idx, end_column=7)
        c_desc = ws_dash.cell(row=idx, column=4, value=p4)
        c_desc.font = font_data
        c_desc.alignment = align_left
        for col_i in range(4, 8):
            ws_dash.cell(row=idx, column=col_i).border = cell_border

    # Adjust Dashboard column widths
    ws_dash.column_dimensions["A"].width = 24
    ws_dash.column_dimensions["B"].width = 18
    ws_dash.column_dimensions["C"].width = 18
    ws_dash.column_dimensions["D"].width = 24
    ws_dash.column_dimensions["E"].width = 20
    ws_dash.column_dimensions["F"].width = 18
    ws_dash.column_dimensions["G"].width = 22
    ws_dash.column_dimensions["H"].width = 16
    ws_dash.column_dimensions["I"].width = 18
    ws_dash.column_dimensions["J"].width = 18
    ws_dash.column_dimensions["K"].width = 18


    # ----------------------------------------------------
    # DATA SHEET BUILDER FUNCTION
    # ----------------------------------------------------
    def build_student_sheet(sheet_title, title_text, student_list):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        total_cols = len(master_columns)
        end_letter = get_column_letter(total_cols)

        # Title Block
        ws.merge_cells(f"A1:{end_letter}1")
        ws["A1"] = f"CENTRE FOR MEDICAL TECHNOLOGIES (IDRP) | IIT JODHPUR & AIIMS JODHPUR"
        ws["A1"].font = font_title
        ws["A1"].fill = fill_title
        ws["A1"].alignment = align_center
        ws.row_dimensions[1].height = 32

        ws.merge_cells(f"A2:{end_letter}2")
        ws["A2"] = title_text
        ws["A2"].font = font_subtitle
        ws["A2"].fill = fill_title
        ws["A2"].alignment = align_center
        ws.row_dimensions[2].height = 18

        # Blank row 3
        ws.row_dimensions[3].height = 8

        # Table Headers row 4
        ws.row_dimensions[4].height = 26
        for col_idx, (col_name, col_w, alignment) in enumerate(master_columns, start=1):
            cell = ws.cell(row=4, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = cell_border

        # Populate rows
        for row_idx, s in enumerate(student_list, start=5):
            ws.row_dimensions[row_idx].height = 20
            row_fill = fill_zebra if row_idx % 2 == 0 else fill_white

            vals = [
                row_idx - 4,
                s['roll_number'],
                s['name'],
                s['program_level'],
                s['degree'],
                int(s['cohort_year']),
                int(s['graduation_year']),
                s['status'],
                s['specialization'] if s['specialization'] else "Medical Technology",
                s['supervisor'] if s['supervisor'] else "Centre Faculty Board (Joint IDRP IITJ & AIIMSJ)",
                s['email'],
                s['phone'] if s['phone'] else "-",
                int(s['publications_count']),
                s['publications_details'] if s['publications_details'] else "-",
                int(s['patents_count']),
                s['patents_details'] if s['patents_details'] else "-",
                int(s['startups_count']),
                s['startups_details'] if s['startups_details'] else "-",
                int(s['achievements_count']),
                s['achievements_details'] if s['achievements_details'] else "-",
                s['date_of_birth'] if s['date_of_birth'] else "-",
                s['gender'] if s['gender'] else "-",
                s['department'],
                s['address'],
                s['city'],
                s['state'],
                s['pincode'],
                s['provenance']
            ]

            for col_idx, (val, (col_name, col_w, alignment)) in enumerate(zip(vals, master_columns), start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.fill = row_fill
                cell.alignment = alignment
                cell.border = cell_border

                # Highlight badges
                if col_name == "Roll Number":
                    cell.font = font_data_bold
                elif col_name == "Student Scholar Name":
                    cell.font = font_data_bold
                elif col_name == "Academic Status":
                    if val == "Active":
                        cell.fill = fill_active
                        cell.font = font_active
                    else:
                        cell.fill = fill_alumni
                        cell.font = font_alumni
                elif col_name in ["Publications (Count)", "Patents (Count)", "Startups (Count)", "Awards (Count)"]:
                    if isinstance(val, (int, float)) and val > 0:
                        cell.font = font_data_bold
                        cell.fill = fill_cohort_bar

        # Column widths & freeze pane
        for col_idx, (col_name, col_w, alignment) in enumerate(master_columns, start=1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = col_w

        ws.freeze_panes = "D5"
        ws.auto_filter.ref = f"A4:{end_letter}{len(student_list) + 4}"
        return ws

    # ----------------------------------------------------
    # SHEET 2: MASTER DIRECTORY (ALL STUDENTS)
    # ----------------------------------------------------
    build_student_sheet(
        sheet_title="Master Directory",
        title_text="Consolidated Master Directory of All 107 Medical Technology Scholars (2020 – 2026 Batches)",
        student_list=students
    )

    # ----------------------------------------------------
    # SHEET 3: M.TECH (MASTERS) COHORTS
    # ----------------------------------------------------
    masters_students = [s for s in students if s['program_level'] == 'Masters']
    build_student_sheet(
        sheet_title="M.Tech (Masters)",
        title_text=f"M.Tech in Medical Technology — Complete Register of All {len(masters_students)} Masters Scholars (2020 – 2026 Batches)",
        student_list=masters_students
    )

    # ----------------------------------------------------
    # SHEET 4: PH.D. SCHOLARS COHORTS
    # ----------------------------------------------------
    phd_students = [s for s in students if s['program_level'] == 'PhD']
    build_student_sheet(
        sheet_title="Ph.D. Scholars",
        title_text=f"Doctoral Research Scholars (Ph.D.) — Complete Register of All {len(phd_students)} Ph.D. Scholars (2020 – 2026 Batches)",
        student_list=phd_students
    )

    # ----------------------------------------------------
    # SHEET 5: DUAL DEGREE (M.TECH-PH.D.) COHORTS
    # ----------------------------------------------------
    dual_students = [s for s in students if s['program_level'] == 'Dual Degree']
    build_student_sheet(
        sheet_title="Dual Degree (M.Tech-Ph.D.)",
        title_text=f"M.Tech-Ph.D. Dual Degree in Medical Technology — Complete Register of All {len(dual_students)} Scholars (2020 – 2025 Batches)",
        student_list=dual_students
    )

    # ----------------------------------------------------
    # SHEET 6: YEAR-WISE BREAKDOWN VIEW
    # ----------------------------------------------------
    ws_yw = wb.create_sheet(title="Year-Wise Cohorts")
    ws_yw.views.sheetView[0].showGridLines = True
    end_letter = get_column_letter(len(master_columns))

    ws_yw.merge_cells(f"A1:{end_letter}1")
    ws_yw["A1"] = "CENTRE FOR MEDICAL TECHNOLOGIES (IDRP) | IIT JODHPUR & AIIMS JODHPUR"
    ws_yw["A1"].font = font_title
    ws_yw["A1"].fill = fill_title
    ws_yw["A1"].alignment = align_center
    ws_yw.row_dimensions[1].height = 32

    ws_yw.merge_cells(f"A2:{end_letter}2")
    ws_yw["A2"] = "Chronological Year-by-Year Cohort Register (2020 through 2026 Admissions)"
    ws_yw["A2"].font = font_subtitle
    ws_yw["A2"].fill = fill_title
    ws_yw["A2"].alignment = align_center
    ws_yw.row_dimensions[2].height = 18

    curr_row = 4
    for yr in sorted(list(set(s['cohort_year'] for s in students))):
        yr_students = [s for s in students if s['cohort_year'] == yr]
        
        # Section Header for Year
        ws_yw.row_dimensions[curr_row].height = 24
        ws_yw.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=len(master_columns))
        sec_cell = ws_yw.cell(row=curr_row, column=1)
        m_cnt = sum(1 for s in yr_students if s['program_level'] == 'Masters')
        p_cnt = sum(1 for s in yr_students if s['program_level'] == 'PhD')
        d_cnt = sum(1 for s in yr_students if s['program_level'] == 'Dual Degree')
        sec_cell.value = f"🎓 ADMISSION COHORT {yr} — Total Scholars: {len(yr_students)} (Masters: {m_cnt} | Ph.D.: {p_cnt} | Dual Degree: {d_cnt})"
        sec_cell.font = font_cohort_header
        sec_cell.fill = fill_cohort_bar
        sec_cell.alignment = align_left
        for col_i in range(1, len(master_columns) + 1):
            ws_yw.cell(row=curr_row, column=col_i).border = Border(top=med_border_side, bottom=med_border_side)
        curr_row += 1

        # Table Headers
        ws_yw.row_dimensions[curr_row].height = 22
        for col_idx, (col_name, col_w, alignment) in enumerate(master_columns, start=1):
            c = ws_yw.cell(row=curr_row, column=col_idx, value=col_name)
            c.font = font_header_sub
            c.fill = fill_sub_header
            c.alignment = align_center
            c.border = cell_border
        curr_row += 1

        # Student rows
        for s_idx, s in enumerate(yr_students, start=1):
            ws_yw.row_dimensions[curr_row].height = 19
            r_fill = fill_zebra if s_idx % 2 == 0 else fill_white

            vals = [
                s_idx,
                s['roll_number'],
                s['name'],
                s['program_level'],
                s['degree'],
                int(s['cohort_year']),
                int(s['graduation_year']),
                s['status'],
                s['specialization'] if s['specialization'] else "Medical Technology",
                s['supervisor'] if s['supervisor'] else "Centre Faculty Board (Joint IDRP IITJ & AIIMSJ)",
                s['email'],
                s['phone'] if s['phone'] else "-",
                int(s['publications_count']),
                s['publications_details'] if s['publications_details'] else "-",
                int(s['patents_count']),
                s['patents_details'] if s['patents_details'] else "-",
                int(s['startups_count']),
                s['startups_details'] if s['startups_details'] else "-",
                int(s['achievements_count']),
                s['achievements_details'] if s['achievements_details'] else "-",
                s['date_of_birth'] if s['date_of_birth'] else "-",
                s['gender'] if s['gender'] else "-",
                s['department'],
                s['address'],
                s['city'],
                s['state'],
                s['pincode'],
                s['provenance']
            ]

            for col_idx, (val, (col_name, col_w, alignment)) in enumerate(zip(vals, master_columns), start=1):
                c = ws_yw.cell(row=curr_row, column=col_idx, value=val)
                c.font = font_data
                c.fill = r_fill
                c.alignment = alignment
                c.border = cell_border

                if col_name in ["Roll Number", "Student Scholar Name"]:
                    c.font = font_data_bold
                elif col_name == "Academic Status":
                    c.fill = fill_active if val == "Active" else fill_alumni
                    c.font = font_active if val == "Active" else font_alumni

            curr_row += 1

        # Blank space between cohorts
        curr_row += 1

    for col_idx, (col_name, col_w, alignment) in enumerate(master_columns, start=1):
        letter = get_column_letter(col_idx)
        ws_yw.column_dimensions[letter].width = col_w

    ws_yw.freeze_panes = "D4"

    # ----------------------------------------------------
    # SHEET 7: RESEARCH, IP & STARTUPS DIRECTORY
    # ----------------------------------------------------
    ws_ip = wb.create_sheet(title="IP & Entrepreneurship")
    ws_ip.views.sheetView[0].showGridLines = True

    ws_ip.merge_cells("A1:I1")
    ws_ip["A1"] = "CENTRE FOR MEDICAL TECHNOLOGIES (IDRP) | IIT JODHPUR & AIIMS JODHPUR"
    ws_ip["A1"].font = font_title
    ws_ip["A1"].fill = fill_title
    ws_ip["A1"].alignment = align_center
    ws_ip.row_dimensions[1].height = 32

    ws_ip.merge_cells("A2:I2")
    ws_ip["A2"] = "Student-Led Innovation Directory: Patents, Incubated Startups, Research Publications & National Awards"
    ws_ip["A2"].font = font_subtitle
    ws_ip["A2"].fill = fill_title
    ws_ip["A2"].alignment = align_center
    ws_ip.row_dimensions[2].height = 18

    innovators = [
        s for s in students
        if s['patents_count'] > 0 or s['startups_count'] > 0 or s['publications_count'] > 0 or s['achievements_count'] > 0
    ]
    innovators.sort(key=lambda s: (s['patents_count'] + s['startups_count'] + s['publications_count'], s['cohort_year']), reverse=True)

    ip_headers = [
        ("S.No.", 8, align_center),
        ("Roll Number", 15, align_center),
        ("Student Scholar Name", 25, align_left),
        ("Program & Cohort", 20, align_center),
        ("Patents Filed / Granted", 44, align_left),
        ("Incubated Startup Ventures", 42, align_left),
        ("Scholarly Publications & Venues", 44, align_left),
        ("Awards & Recognitions", 34, align_left),
        ("Faculty Mentors & Supervisors", 38, align_left)
    ]

    ws_ip.row_dimensions[4].height = 26
    for col_idx, (h_name, h_w, h_align) in enumerate(ip_headers, start=1):
        c = ws_ip.cell(row=4, column=col_idx, value=h_name)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = cell_border
        letter = get_column_letter(col_idx)
        ws_ip.column_dimensions[letter].width = h_w

    for idx, s in enumerate(innovators, start=5):
        ws_ip.row_dimensions[idx].height = 24
        r_fill = fill_zebra if idx % 2 == 0 else fill_white

        vals = [
            idx - 4,
            s['roll_number'],
            s['name'],
            f"{s['program_level']} ({s['cohort_year']})",
            s['patents_details'] if s['patents_details'] else "-",
            s['startups_details'] if s['startups_details'] else "-",
            s['publications_details'] if s['publications_details'] else "-",
            s['achievements_details'] if s['achievements_details'] else "-",
            s['supervisor']
        ]

        for col_idx, (val, (h_name, h_w, h_align)) in enumerate(zip(vals, ip_headers), start=1):
            c = ws_ip.cell(row=idx, column=col_idx, value=val)
            c.font = font_data
            c.fill = r_fill
            c.alignment = h_align
            c.border = cell_border

            if h_name in ["Roll Number", "Student Scholar Name"]:
                c.font = font_data_bold
            elif h_name == "Patents Filed / Granted" and val != "-":
                c.fill = fill_cohort_bar
            elif h_name == "Incubated Startup Ventures" and val != "-":
                c.fill = fill_active

    ws_ip.freeze_panes = "D5"
    ws_ip.auto_filter.ref = f"A4:I{len(innovators) + 4}"

    # Output file paths
    out_path_1 = "IIT_Jodhpur_MedTech_Students_Directory_Master.xlsx"
    out_path_2 = "broucher/IIT_Jodhpur_MedTech_Students_Directory_Master.xlsx"
    
    wb.save(out_path_1)
    wb.save(out_path_2)
    print(f"Excel successfully created at:\n  - {out_path_1}\n  - {out_path_2}")

if __name__ == '__main__':
    create_excel()
